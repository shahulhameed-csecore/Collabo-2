import io
import logging
from typing import Dict, Any, List
from datetime import datetime
from collections import defaultdict
import openpyxl
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from app.core.config import settings
from app.services.supabase import get_supabase_admin

logger = logging.getLogger(__name__)

# get_service_client removed, using get_supabase_admin from app.services.supabase

async def fetch_monthly_metrics(user_id: str, start_date: str, end_date: str) -> Dict[str, Any]:
    """
    Fetch and aggregate metrics for the given user and date range.
    start_date and end_date should be in 'YYYY-MM-DD' format.
    """
    client = await get_supabase_admin()
    
    # Query all campaigns for this user created in the month range
    # Assuming we filter by created_at or deadline. Let's use created_at.
    import asyncio
    resp = await (client.table("campaigns").select(
        "status, payment_amount, platform"
    ).eq("user_id", user_id).gte("created_at", f"{start_date}T00:00:00Z").lte("created_at", f"{end_date}T23:59:59Z").execute())
    
    data = resp.data or []
    
    total_spend = 0.0
    paid_count = 0
    cancelled_count = 0
    pending_count = 0
    pending_value = 0.0
    platform_spend = defaultdict(float)
    
    for c in data:
        status = c.get("status")
        amount = c.get("payment_amount") or 0.0
        platform = c.get("platform") or "Unknown"
        
        if status == "paid":
            total_spend += amount
            paid_count += 1
            platform_spend[platform] += amount
        elif status == "cancelled":
            cancelled_count += 1
        elif status in ["active", "content_received", "approved"]:
            pending_count += 1
            pending_value += amount

    # Success Rate
    total_resolved = paid_count + cancelled_count
    success_rate = (paid_count / total_resolved * 100) if total_resolved > 0 else 0.0
    
    # Top Platforms
    sorted_platforms = sorted(platform_spend.items(), key=lambda x: x[1], reverse=True)
    top_5_platforms = [{"platform": p[0], "spend": p[1]} for p in sorted_platforms[:5]]
    
    return {
        "total_spend": total_spend,
        "success_rate": success_rate,
        "pending_count": pending_count,
        "pending_value": pending_value,
        "top_platforms": top_5_platforms,
        "total_campaigns_in_period": len(data)
    }

def generate_pdf_report(metrics: Dict[str, Any], month_str: str) -> bytes:
    """
    Generate a PDF report using ReportLab.
    """
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter)
    styles = getSampleStyleSheet()
    elements = []
    
    # Title
    title_style = styles['Heading1']
    elements.append(Paragraph(f"Monthly Campaign Report - {month_str}", title_style))
    elements.append(Spacer(1, 20))
    
    # Summary Metrics Table
    summary_data = [
        ["Metric", "Value"],
        ["Total Spend", f"INR {metrics['total_spend']:,.2f}"],
        ["Success Rate", f"{metrics['success_rate']:.1f}%"],
        ["Pending Payments Count", str(metrics['pending_count'])],
        ["Pending Payments Value", f"INR {metrics['pending_value']:,.2f}"],
        ["Total Campaigns Started", str(metrics['total_campaigns_in_period'])]
    ]
    
    t_summary = Table(summary_data, colWidths=[250, 150])
    t_summary.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#10b981")),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor("#f8fafc")),
        ('GRID', (0, 0), (-1, -1), 1, colors.HexColor("#cbd5e1"))
    ]))
    elements.append(t_summary)
    elements.append(Spacer(1, 30))
    
    # Top Platforms Table
    elements.append(Paragraph("Top Platforms by Spend", styles['Heading2']))
    elements.append(Spacer(1, 10))
    
    platform_data = [["Platform", "Spend (INR)"]]
    if metrics["top_platforms"]:
        for p in metrics["top_platforms"]:
            platform_data.append([p["platform"], f"INR {p['spend']:,.2f}"])
    else:
        platform_data.append(["No paid data", "-"] )

    t_plat = Table(platform_data, colWidths=[250, 150])
    t_plat.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#3b82f6")),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor("#f8fafc")),
        ('GRID', (0, 0), (-1, -1), 1, colors.HexColor("#cbd5e1"))
    ]))
    elements.append(t_plat)
    
    doc.build(elements)
    pdf_bytes = buffer.getvalue()
    buffer.close()
    return pdf_bytes

def generate_excel_report(metrics: Dict[str, Any], month_str: str) -> bytes:
    """
    Generate an Excel report using openpyxl.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Monthly Summary"
    
    # Title
    ws["A1"] = f"Monthly Campaign Report - {month_str}"
    ws["A1"].font = openpyxl.styles.Font(size=14, bold=True)
    
    # Headers
    headers = ["Metric", "Value"]
    ws.append([]) # Empty row
    ws.append(headers)
    
    # Style headers
    for col in range(1, 3):
        cell = ws.cell(row=3, column=col)
        cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
        cell.fill = openpyxl.styles.PatternFill(start_color="10B981", end_color="10B981", fill_type="solid")
        
    # Summary Data
    ws.append(["Total Spend (INR)", metrics['total_spend']])
    ws.append(["Success Rate (%)", round(metrics['success_rate'], 2)])
    ws.append(["Pending Payments Count", metrics['pending_count']])
    ws.append(["Pending Payments Value (INR)", metrics['pending_value']])
    ws.append(["Total Campaigns Started", metrics['total_campaigns_in_period']])
    
    # Top Platforms
    ws.append([])
    ws.append(["Top Platforms by Spend", "Spend (INR)"])
    row = ws.max_row
    for col in range(1, 3):
        cell = ws.cell(row=row, column=col)
        cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
        cell.fill = openpyxl.styles.PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
        
    if metrics["top_platforms"]:
        for p in metrics["top_platforms"]:
            ws.append([p["platform"], p["spend"]])
    else:
        ws.append(["No paid data", 0])
        
    # Formatting widths
    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 20
    
    buffer = io.BytesIO()
    wb.save(buffer)
    excel_bytes = buffer.getvalue()
    buffer.close()
    return excel_bytes
